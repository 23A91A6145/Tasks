"""
Hashira ATS - Robust Document Parser
Supports PDF (via PyMuPDF), DOCX (via python-docx), and plain text files.
Handles both disk paths and Streamlit UploadedFile (BytesIO) streams safely.
"""
import io
import re
from typing import Tuple, Optional

def extract_text_from_pdf(file_bytes_or_path) -> str:
    """Extract text from PDF using PyMuPDF (fitz)."""
    import fitz  # PyMuPDF
    text_chunks = []
    if isinstance(file_bytes_or_path, (str, bytes, bytearray)):
        if isinstance(file_bytes_or_path, str):
            doc = fitz.open(file_bytes_or_path)
        else:
            doc = fitz.open(stream=file_bytes_or_path, filetype="pdf")
    else:
        # Streamlit UploadedFile or BytesIO
        content = file_bytes_or_path.read()
        if hasattr(file_bytes_or_path, "seek"):
            file_bytes_or_path.seek(0)
        doc = fitz.open(stream=content, filetype="pdf")

    for page_num in range(len(doc)):
        page = doc[page_num]
        text_chunks.append(page.get_text())
    doc.close()
    return "\n".join(text_chunks)

def extract_text_from_docx(file_bytes_or_path) -> str:
    """Extract text from DOCX using python-docx."""
    from docx import Document
    if isinstance(file_bytes_or_path, str):
        doc = Document(file_bytes_or_path)
    elif isinstance(file_bytes_or_path, (bytes, bytearray)):
        doc = Document(io.BytesIO(file_bytes_or_path))
    else:
        content = file_bytes_or_path.read()
        if hasattr(file_bytes_or_path, "seek"):
            file_bytes_or_path.seek(0)
        doc = Document(io.BytesIO(content))

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(paragraphs)

def extract_text_from_txt(file_bytes_or_path) -> str:
    """Extract text from plain text input with encoding fallback."""
    if isinstance(file_bytes_or_path, str):
        with open(file_bytes_or_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    elif isinstance(file_bytes_or_path, (bytes, bytearray)):
        try:
            return file_bytes_or_path.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes_or_path.decode("latin-1", errors="ignore")
    else:
        content = file_bytes_or_path.read()
        if hasattr(file_bytes_or_path, "seek"):
            file_bytes_or_path.seek(0)
        if isinstance(content, str):
            return content
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError:
            return content.decode("latin-1", errors="ignore")

def extract_text(file_obj, filename: str) -> Tuple[bool, str, Optional[str]]:
    """
    Safely extract text from any supported file.
    Returns: (success: bool, extracted_text: str, error_message: Optional[str])
    """
    lower_name = filename.lower()
    try:
        if lower_name.endswith(".pdf"):
            text = extract_text_from_pdf(file_obj)
        elif lower_name.endswith(".docx"):
            text = extract_text_from_docx(file_obj)
        elif lower_name.endswith((".txt", ".md")):
            text = extract_text_from_txt(file_obj)
        else:
            return False, "", f"Unsupported file extension for '{filename}'. Allowed: PDF, DOCX, TXT."
        
        cleaned = clean_extracted_text(text)
        if not cleaned or len(cleaned.strip()) < 20:
            return False, "", f"Unable to extract meaningful text from '{filename}'. Document may be empty or scanned images."
        
        return True, cleaned, None
    except Exception as e:
        return False, "", f"Error parsing '{filename}': {str(e)}"

def clean_extracted_text(text: str) -> str:
    """Standardize whitespace and remove non-printable characters."""
    if not text:
        return ""
    # Replace weird unicode spaces
    text = text.replace('\xa0', ' ').replace('\u200b', '')
    # Condense consecutive blank lines
    text = re.sub(r'\n{3,}', '\n\n', text)
    # Condense multiple horizontal spaces
    text = re.sub(r'[ \t]{2,}', ' ', text)
    return text.strip()

def detect_candidate_name(text: str, fallback_filename: str) -> str:
    """
    Extract candidate name from header of resume or fallback to clean filename.
    """
    lines = [l.strip() for l in text.split('\n') if l.strip()]
    for line in lines[:6]:
        # Filter out common headers or contact details
        if re.search(r'@|github\.com|linkedin\.com|phone|\+91|\d{10}|resume|curriculum', line, re.IGNORECASE):
            continue
        words = line.split()
        if 1 <= len(words) <= 4 and all(w.isalpha() or w in ['.', '-'] for w in words):
            return line.title()
    
    # Fallback to filename without extension
    clean_name = re.sub(r'[\-_]', ' ', fallback_filename.rsplit('.', 1)[0])
    clean_name = re.sub(r'(?i)resume|cv|latest|final|updated|\d+', '', clean_name).strip()
    return clean_name.title() if clean_name else "Candidate"
